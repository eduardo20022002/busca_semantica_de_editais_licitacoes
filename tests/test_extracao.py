import io
import zipfile

from docx import Document
from openpyxl import Workbook

from editais.extracao import TextoExtraido, extrair_textos_de_arquivo


def _pdf_com_texto(texto: str) -> bytes:
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = b"BT /F1 24 Tf 100 700 Td (" + texto.encode("latin-1") + b") Tj ET"
    objetos.append(
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream"
    )
    objetos.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    pdf = b"%PDF-1.4\n"
    offsets = []
    for indice, obj in enumerate(objetos, start=1):
        offsets.append(len(pdf))
        pdf += str(indice).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    posicao_xref = len(pdf)
    pdf += b"xref\n0 " + str(len(objetos) + 1).encode() + b"\n"
    pdf += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf += ("%010d 00000 n \n" % offset).encode()
    pdf += b"trailer\n<< /Size " + str(len(objetos) + 1).encode() + b" /Root 1 0 R >>\n"
    pdf += b"startxref\n" + str(posicao_xref).encode() + b"\n%%EOF"
    return pdf


def _docx_com_texto(*paragrafos: str) -> bytes:
    documento = Document()
    for paragrafo in paragrafos:
        documento.add_paragraph(paragrafo)
    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def _xlsx_com_celulas(linhas: list[list[str]]) -> bytes:
    livro = Workbook()
    planilha = livro.active
    assert planilha is not None
    for linha in linhas:
        planilha.append(linha)
    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def _zip_com(arquivos: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zip_file:
        for nome, conteudo in arquivos.items():
            zip_file.writestr(nome, conteudo)
    return buffer.getvalue()


def test_extrai_texto_de_txt() -> None:
    resultado = extrair_textos_de_arquivo("nota.txt", "objeto do edital".encode("utf-8"))

    assert resultado == [TextoExtraido(nome_arquivo="nota.txt", texto="objeto do edital")]


def test_extrai_texto_de_pdf() -> None:
    conteudo = _pdf_com_texto("ESPECIFICACOES TECNICAS do sistema")

    resultado = extrair_textos_de_arquivo("edital.pdf", conteudo)

    assert len(resultado) == 1
    assert resultado[0].nome_arquivo == "edital.pdf"
    assert "ESPECIFICACOES TECNICAS" in resultado[0].texto


def test_extrai_texto_de_docx() -> None:
    conteudo = _docx_com_texto("DO OBJETO", "Contratacao de software de gestao")

    resultado = extrair_textos_de_arquivo("termo.docx", conteudo)

    assert len(resultado) == 1
    assert "DO OBJETO" in resultado[0].texto
    assert "software de gestao" in resultado[0].texto


def test_extrai_texto_de_xlsx() -> None:
    conteudo = _xlsx_com_celulas([["Item", "Descricao"], ["1", "Licenca de telemedicina"]])

    resultado = extrair_textos_de_arquivo("planilha.xlsx", conteudo)

    assert len(resultado) == 1
    assert "Licenca de telemedicina" in resultado[0].texto


def test_extrai_recursivamente_de_zip() -> None:
    conteudo = _zip_com(
        {
            "a.txt": "primeiro arquivo".encode("utf-8"),
            "b.txt": "segundo arquivo".encode("utf-8"),
        }
    )

    resultado = extrair_textos_de_arquivo("anexos.zip", conteudo)

    textos = {t.texto for t in resultado}
    assert textos == {"primeiro arquivo", "segundo arquivo"}


def test_extrai_zip_dentro_de_zip() -> None:
    interno = _zip_com({"dentro.txt": "texto aninhado".encode("utf-8")})
    externo = _zip_com({"nested.zip": interno, "fora.txt": "texto de fora".encode("utf-8")})

    resultado = extrair_textos_de_arquivo("pacote.zip", externo)

    textos = {t.texto for t in resultado}
    assert textos == {"texto aninhado", "texto de fora"}


def test_formato_nao_suportado_retorna_vazio() -> None:
    assert extrair_textos_de_arquivo("imagem.png", b"\x89PNG\r\n") == []
    assert extrair_textos_de_arquivo("compactado.rar", b"Rar!\x1a\x07") == []


def test_arquivo_corrompido_retorna_vazio_sem_levantar() -> None:
    assert extrair_textos_de_arquivo("quebrado.pdf", b"nao e um pdf de verdade") == []


def test_texto_vazio_ou_so_espacos_retorna_vazio() -> None:
    assert extrair_textos_de_arquivo("vazio.txt", b"   \n  \t ") == []
